import os
import pandas as pd
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfbase import pdfmetrics
# txtファイルのパス
txt_path = '/Users/kamehaku/Desktop/Add-LLM vs Chat-GPT/tools/DataSet/TechJa.txt'
# DataSetフォルダのパス
dataset_path = '/Users/kamehaku/Desktop/Add-LLM vs Chat-GPT/tools/DataSet/'

# txtファイルを読み込む
df = pd.read_csv(txt_path, sep='\t', dtype=str)

# 出力ファイル名をtxtファイル名と同じにする
base_name = os.path.splitext(os.path.basename(txt_path))[0]

# DataFrameをdocxファイルとして保存
docx_filename = os.path.join(dataset_path, f'{base_name}.docx')
doc = Document()
table = doc.add_table(rows=1, cols=len(df.columns))
table.style = 'Table Grid'

# ヘッダーを追加
for i, col_name in enumerate(df.columns):
    table.cell(0, i).text = str(col_name)  # col_nameを文字列に変換

# データを行ごとに追加
for row in df.itertuples(index=False):
    cells = table.add_row().cells
    for i, value in enumerate(row):
        cells[i].text = str(value)  # valueを文字列に変換

doc.save(docx_filename)

# フォントファイルのフルパスを指定
font_path = '/Users/kamehaku/Desktop/Add-LLM vs Chat-GPT/tools/Noto_Sans_JP/static/NotoSansJP-Regular.ttf'

# フォントの登録
pdfmetrics.registerFont(TTFont('NotoSansJP', font_path))

# PDFファイルの作成
pdf_file = os.path.join(dataset_path, f"{base_name}.pdf")
doc = SimpleDocTemplate(pdf_file, pagesize=letter)

# スタイルの設定
styles = getSampleStyleSheet()
notosans_style = styles['Normal'].clone('NotoSansJP')  # 既存のスタイルをコピー
notosans_style.fontName = 'NotoSansJP'
notosans_style.fontSize = 10
styles.add(notosans_style)  # 新しいスタイルを追加

# データをParagraphオブジェクトに変換
elements = []
for col_name in df.columns:
    elements.append(Paragraph(str(col_name), styles['NotoSansJP']))
for row in df.values.tolist():
    for value in row:
        elements.append(Paragraph(str(value), styles['NotoSansJP']))

# PDFを生成
doc.build(elements)