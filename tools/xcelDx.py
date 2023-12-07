import os
import pandas as pd
from docx import Document
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter

# Excelファイルのパス
xls_path = '/Users/kamehaku/Desktop/Add-LLM vs Chat-GPT/tools/DataX/aj10.xls'
# DataSetフォルダのパス
dataset_path = '/Users/kamehaku/Desktop/Add-LLM vs Chat-GPT/tools/DataSet/'

# Excelファイルを読み込む（数値データを文字列として解析）
df = pd.read_excel(xls_path, dtype=str)

# DataFrameをtxtファイルとして保存
txt_filename = os.path.join(dataset_path, 'aj10.txt')
with open(txt_filename, 'w', encoding='utf-8') as text_file:
    text_file.write(df.to_csv(sep='\t', index=False))

# DataFrameをdocxファイルとして保存
docx_filename = os.path.join(dataset_path, 'aj10.docx')
doc = Document()
table = doc.add_table(rows=1, cols=len(df.columns))
table.style = 'Table Grid'

# ヘッダーを追加
for i, col_name in enumerate(df.columns):
    table.cell(0, i).text = str(col_name)  # col_nameを文字列に変換

# データを行ごとに追加
for row in df.itertuples(index=False):
    cells = table.add_row().cells
    for i, value in enumerate(row):
        cells[i].text = str(value)  # valueを文字列に変換

doc.save(docx_filename)

# DataFrameをPDFファイルとして保存
pdf_filename = os.path.join(dataset_path, 'aj10.pdf')
c = canvas.Canvas(pdf_filename, pagesize=letter)
y = 800  # 初期のY座標

# ヘッダーを追加
for col_name in df.columns:
    c.drawString(72, y, str(col_name))  # col_nameを文字列に変換
    y -= 12

# データを行ごとに追加
for row in df.values.tolist():
    for value in row:
        c.drawString(72, y, str(value))  # valueを文字列に変換
    y -= 12
    if y < 72:  # ページが終わったら新しいページを開始
        c.showPage()
        y = 800

c.save()